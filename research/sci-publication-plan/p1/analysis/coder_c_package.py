"""Build the blind package for human coder C.
One randomly chosen image per P1/P2 cell (48) plus, if a P3 folder is given, one per channel x role (24),
renamed H01.. in random order. The key stays in data/coder_C_key.csv.
Usage: python coder_c_package.py <P1/P2 image dir> <out dir> [<P3 image dir>]"""
import sys, csv, random, shutil, subprocess, pathlib
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

src, out = pathlib.Path(sys.argv[1]), pathlib.Path(sys.argv[2])
p3 = pathlib.Path(sys.argv[3]) if len(sys.argv) > 3 else None
P = pathlib.Path(__file__).resolve().parents[1]
rng = random.Random(20261022)

key = list(csv.DictReader(open(P / 'data/sheet_key.csv')))
cells = {}
for r in key:
    cells.setdefault((r['model'], r['strategy'], r['role']), []).append(r)
pick = []
for _, v in sorted(cells.items()):
    r = rng.choice(v)
    pick.append({'img': f"{r['sheet']}-{r['rep']}", 'model': r['model'], 'strategy': r['strategy'],
                 'role': r['role'], 'rep': r['rep'], 'path': next(src.rglob(r['file']))})
if p3:
    files = {}
    for f in sorted(p3.rglob('*_P3_*.jpg')):
        m, _, role, rep = f.stem.split('_')
        files.setdefault((m, role), []).append((rep, f))
    for (m, role), v in sorted(files.items()):
        rep, f = rng.choice(v)
        pick.append({'img': f.stem, 'model': m, 'strategy': 'P3', 'role': role, 'rep': rep, 'path': f})
rng.shuffle(pick)

roles = {'R1': ('诸葛亮', '空城计'), 'R2': ('关羽', '单刀会'), 'R3': ('杨玉环', '贵妃醉酒'),
         'R4': ('穆桂英', '穆桂英挂帅'), 'R5': ('包拯', '铡美案'), 'R6': ('赵云', '长坂坡'),
         'R7': ('苏三', '玉堂春·起解'), 'R8': ('蒋干', '群英会')}
(out / 'images').mkdir(parents=True, exist_ok=True)
rows = []
with open(P / 'data/coder_C_key.csv', 'w', newline='') as fk:
    w = csv.writer(fk); w.writerow(['hid', 'img', 'model', 'strategy', 'role', 'rep', 'file'])
    for i, r in enumerate(pick, 1):
        hid = f'H{i:02d}'
        shutil.copy(r['path'], out / 'images' / f"{hid}_{r['role']}.jpg")
        w.writerow([hid, r['img'], r['model'], r['strategy'], r['role'], r['rep'], r['path'].name])
        rows.append((hid, r['role'], *roles[r['role']]))

wb = Workbook(); ws = wb.active; ws.title = '编码'
ws.append(['编号', '角色ID', '角色', '剧目', 'D1形制', 'D2色彩', 'D3纹样', 'D4头戴', 'D5配件与标志', 'D6跨文化', '备注（可选）'])
for row in rows:
    ws.append(list(row))
dv = DataValidation(type='list', formula1='"0,1,9"', allow_blank=True); ws.add_data_validation(dv)
dv.add(f'E2:J{len(rows) + 1}')
for c, wd in zip('ABCDEFGHIJK', [6, 7, 8, 13, 8, 8, 8, 8, 12, 9, 40]):
    ws.column_dimensions[c].width = wd

items = {
    'R1': ['主服：八卦衣', '头戴：八卦巾', '配件：羽扇', '髯口：黑三髯；俊扮'],
    'R2': ['主服：绿蟒（绿靠也算符合）', '主色：绿', '头戴：夫子盔，不插翎子', '配件：玉带', '脸与髯：红脸、黑长髯'],
    'R3': ['主服：女蟒（及膝，衬裙）', '头戴：凤冠', '配件：云肩、玉带、水袖', '俊扮无髯'],
    'R4': ['主服：女靠', '头戴：额子插两根长翎子', '配件：靠旗', '俊扮'],
    'R5': ['主服：黑蟒', '主色：黑', '头戴：相貂', '配件：玉带', '脸与髯：黑整脸、额白月牙、黑满髯'],
    'R6': ['主服：白靠', '主色：白', '头戴：武将盔头', '配件：靠旗、枪', '俊扮无髯'],
    'R7': ['主服：红罪衣罪裙', '主色：红', '头戴：包头，不戴冠', '配件：鱼形枷、水袖', '俊扮'],
    'R8': ['主服：褶子', '头戴：方巾', '脸：豆腐块丑脸'],
}
rv = wb.create_sheet('规制审读')
rv.append(['角色ID', '角色（剧目）', '规制要点', '判断（同意/不同意/流派差异）', '说明'])
for rid, its in items.items():
    for it in its:
        rv.append([rid, f'{roles[rid][0]}（{roles[rid][1]}）', it])
for d in ['D1 只看主服大类，做工不算', 'D2 只对关羽、包拯、赵云、苏三编码；其余仅越制记违规',
          'D3 只看纹样种类，不看画工', 'D5 替代品不算符合', 'D6 清宫、日韩、现代汉服、西式元素记违规']:
    rv.append(['规则', '', d])
for _ in range(5):
    rv.append(['补充', '', '（规制表没写、但舞台上会被看作穿错的规矩）'])
dv2 = DataValidation(type='list', formula1='"同意,不同意,流派差异"', allow_blank=True); rv.add_data_validation(dv2)
dv2.add(f'D2:D{rv.max_row}')
for c, wd in zip('ABCDE', [8, 22, 40, 26, 50]):
    rv.column_dimensions[c].width = wd
wb.save(out / '编码表_编码者C.xlsx')

md = (P / 'coder_C_instructions.md').read_text(encoding='utf-8').replace('{N}', str(len(rows)))
tmp = out / '编码说明.md'; tmp.write_text(md, encoding='utf-8')
subprocess.run(['pandoc', str(tmp), '-o', str(out / '编码说明.docx')], check=True); tmp.unlink()
print(len(rows), 'images')

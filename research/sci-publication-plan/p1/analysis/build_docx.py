"""Build JTI submission files from manuscript_v2.md.
Tables are moved after the references with 'near here' markers; two versions are written:
full (with title page) and anonymised (author details removed)."""
import re, subprocess, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

P = Path(__file__).resolve().parents[1]
src = (P / 'manuscript_v2.md').read_text(encoding='utf-8')
lines = src.split('\n')

# pull table blocks: caption line, table rows, optional note
tables, out, i = [], [], 0
while i < len(lines):
    m = re.match(r'\*\*Table (\d+)\.\*\*', lines[i])
    if m:
        blk = [lines[i]]; i += 1
        while i < len(lines) and (lines[i].strip() == '' or lines[i].startswith('|')):
            blk.append(lines[i]); i += 1
            if blk[-1].strip() == '' and len(blk) > 2 and blk[-2].startswith('|'):
                break
        if i < len(lines) and lines[i].startswith('*Note.*'):
            blk += [lines[i], '']; i += 1
        tables.append('\n'.join(blk))
        out += [f'[Table {m.group(1)} near here]', '']
        continue
    out.append(lines[i]); i += 1
body = '\n'.join(out)
ins = '\n## Tables\n\n' + '\n\n'.join(tables) + '\n'
body = body.replace('\n---\n\n## Figure captions', ins + '\n## Figure captions')
body = body.replace('\n---\n', '\n')

def anon(t):
    t = re.sub(r'^Zhengxian Yan,.*$', '', t, flags=re.M)
    t = re.sub(r'Coder A is the author, who.*?University of Korea\.',
               'Coder A is the author [background details removed for anonymous review].', t)
    return t

def build(md, name):
    tmp = P / 'submission' / (name + '.md')
    tmp.write_text(md, encoding='utf-8')
    docx = P / 'submission' / (name + '.docx')
    subprocess.run(['pandoc', str(tmp), '-o', str(docx)], check=True)
    tmp.unlink()
    d = Document(docx)
    for st in d.styles:
        if st.type == 1:
            st.font.name = 'Times New Roman'
            st.element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun') if st.element.rPr is not None and st.element.rPr.rFonts is not None else None
            pf = st.paragraph_format
            pf.line_spacing = 2.0
    d.styles['Normal'].font.size = Pt(12)
    for st in d.styles:
        if st.type == 1 and (st.name.startswith('Heading') or st.name == 'Title'):
            st.font.color.rgb = RGBColor(0, 0, 0)
            st.font.size = Pt(14 if st.name == 'Title' else 12)
            if st.element.rPr is not None and st.element.rPr.rFonts is not None:
                for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
                    st.element.rPr.rFonts.attrib.pop(qn(a), None)
            st.font.name = 'Times New Roman'
    for t in d.tables:
        t.style = d.styles['Table Grid'] if 'Table Grid' in [s.name for s in d.styles] else t.style
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
                    for r in p.runs: r.font.size = Pt(9)
    sp = d.sections[0]._sectPr
    ln = OxmlElement('w:lnNumType')
    ln.set(qn('w:countBy'), '1'); ln.set(qn('w:restart'), 'continuous')
    sp.insert(0, ln)
    # page number in footer
    f = d.sections[0].footer.paragraphs[0]
    for kind, txt in (('begin', None), (None, 'PAGE'), ('separate', None), ('end', None)):
        r = f.add_run()
        if kind:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), kind)
        else:
            e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt
        r._r.append(e)
    d.save(docx)
    print('wrote', docx)

build(body, 'manuscript_full')
build(anon(body), 'manuscript_anonymised')
